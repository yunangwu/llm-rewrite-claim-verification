#!/usr/bin/env python3
"""
Build the human-annotation workbook for the semantic-preservation study.

Samples claims from the gpt-oss rewrites, assigns two annotators to every claim,
and writes one .xlsx with a sheet per annotator plus an unblinded master key.

Sampling, per condition, ``--n-per-cond`` claims each. All sampled claims are
mutually disjoint at the claim-pair level: the two members of a SciClaimEval
pair share the same original claim text, so drawing both would show an annotator
the same original twice.

  CT-mix-42, C_1  restricted to the *unflagged* subset of the semantic filter
                  (core.analysis.semantic_filter.flag_reasons). The paper claims
                  the confidence drift persists on rewrites the filter does not
                  flag, so that is the subset human annotation has to validate.
  OA-mix-42, D_1  random. The filter only detects a hedge inserted before a
                  precise element, which does not describe boosting conditions.

Rewrites that are missing or byte-identical to the original are never sampled;
there is nothing for an annotator to judge.

Annotator sheets are blinded: no condition, no claim_id, independently shuffled,
and ``same_as_original`` / ``type`` / ``Comments`` left empty to be filled in.
The condition and claim_id live only in master_key.csv.

Usage:
  python annotation/build_annotation_sheets.py
  python annotation/build_annotation_sheets.py --n-per-cond 80 --seed 20260915
"""

from __future__ import annotations

import argparse
import csv
import difflib
import random
import re
import sys
import unicodedata
from itertools import combinations
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.analysis.semantic_filter import flag_reasons
from core.config import DATASETS
from core.io import load_base_dataset, load_rewrite, rewrite_path

OUT_DIR = Path(__file__).resolve().parent
ANNOTATORS = ["A1", "A2", "A3", "A4", "A5", "A6"]
SPLITS = ("dev", "test")

# (condition, method, type label, restrict to the filter's unflagged subset)
CONDITIONS = [
    ("CT-mix-42", "natural", "CT", True),
    ("C_1", "inject", "C1", True),
    ("D_1", "inject", "D1", False),
    ("OA-mix-42", "natural", "OA", False),
]

# Mirrors the previous annotation file, minus `priority` and the stray blank
# column; `item_id` is added so two annotators' judgments can be merged back.
SHEET_COLUMNS = [
    "item_id", "original_claim", "rewritten_claim",
    "same_as_original", "evi_path", "type", "Comments",
]
COLUMN_WIDTHS = [10, 72, 72, 17, 34, 8, 30]
WRAPPED = {"original_claim", "rewritten_claim"}

DIFF_LEGEND = (
    "Bold red = word added or changed relative to original_claim. "
    "Differences of case or typography only (e.g. hyphen style) are not marked."
)

MASTER_COLUMNS = [
    "item_id", "type", "condition", "method", "split", "claim_id",
    "claim_id_pair", "label", "evi_path", "annotator_1", "annotator_2",
    "unflagged_only", "original_claim", "rewritten_claim",
]


# The rewriters emit typographic variants (non-breaking hyphens, curly quotes)
# that are not edits. Folded away when aligning words so they are not marked;
# the text shown to the annotator is always the stored claim, unmodified.
_TYPOGRAPHY = {
    **dict.fromkeys(map(ord, "\u2010\u2011\u2012\u2013\u2014\u2015\u2212"), "-"),
    **dict.fromkeys(map(ord, "\u2018\u2019\u201b\u2032"), "'"),
    **dict.fromkeys(map(ord, "\u201c\u201d\u2033"), '"'),
    **dict.fromkeys(map(ord, "\u00a0\u2009\u202f\u200b"), " "),
}


def _word_spans(text: str) -> list[tuple[str, int, int]]:
    return [(m.group(), m.start(), m.end()) for m in re.finditer(r"\S+", text)]


def _match_key(word: str) -> str:
    """Comparison key for diff alignment: case- and typography-insensitive."""
    return unicodedata.normalize("NFKC", word.translate(_TYPOGRAPHY)).casefold()


def _mark_runs(
    text: str, spans: list[tuple[str, int, int]], changed: set[int]
) -> list[tuple[str, bool]]:
    """
    Split ``text`` into (segment, marked) runs, marking the words whose index is
    in ``changed``.

    Inter-word whitespace is emitted verbatim rather than re-joining tokens, so
    the rendered cell is byte-identical to the stored claim (11 claims have
    double spaces or non-standard separators). Whitespace between two marked
    words is marked too, so a multi-word edit highlights as one continuous span.
    """
    runs: list[list] = []
    cursor = 0
    for i, (word, start, end) in enumerate(spans):
        if start > cursor:
            runs.append([text[cursor:start], i - 1 in changed and i in changed])
        runs.append([word, i in changed])
        cursor = end
    if cursor < len(text):
        runs.append([text[cursor:], False])

    merged: list[list] = []
    for segment, k in runs:
        if merged and merged[-1][1] == k:
            merged[-1][0] += segment
        else:
            merged.append([segment, k])
    return [(segment, k) for segment, k in merged]


def diff_runs(original: str, rewritten: str) -> list[tuple[str, bool]]:
    """
    Marked runs for ``rewritten``: the words it adds or changes relative to
    ``original``. Only the rewrite side is marked -- the original is shown clean,
    so the annotator reads it as an unannotated reference.
    """
    o_spans, r_spans = _word_spans(original), _word_spans(rewritten)
    matcher = difflib.SequenceMatcher(
        None,
        [_match_key(w) for w, _, _ in o_spans],
        [_match_key(w) for w, _, _ in r_spans],
        autojunk=False,
    )
    added: set[int] = set()
    for tag, _i1, _i2, j1, j2 in matcher.get_opcodes():
        if tag in ("replace", "insert"):
            added.update(range(j1, j2))
    return _mark_runs(rewritten, r_spans, added)


def load_base(dataset: str) -> dict[str, dict]:
    """claim_id -> original claim, evidence path and pair key, over both splits."""
    base: dict[str, dict] = {}
    for split in SPLITS:
        for item in load_base_dataset(DATASETS[dataset], split):
            base[item["claim_id"]] = {
                "claim": (item.get("claim") or "").strip(),
                "evi_path": item.get("evi_path") or "",
                "label": item.get("label") or "",
                "split": split,
                "claim_id_pair": item.get("claim_id_pair") or "",
                "pair_key": f"{split}:{item.get('claim_id_pair')}",
            }
    return base


def build_pairs(base: dict[str, dict]) -> dict[str, list[str]]:
    """pair_key -> the two claim_ids, keeping complete pairs only."""
    pairs: dict[str, list[str]] = {}
    for claim_id, info in base.items():
        pairs.setdefault(info["pair_key"], []).append(claim_id)
    return {k: sorted(v) for k, v in pairs.items() if len(v) == 2}


def eligible_pairs(
    dataset: str,
    condition: str,
    method: str,
    unflagged_only: bool,
    base: dict[str, dict],
    pairs: dict[str, list[str]],
) -> tuple[list[tuple[str, list[str]]], dict[str, int]]:
    """Pairs with at least one annotatable member, plus exclusion counts."""
    rewrites: dict[str, str] = {}
    for split in SPLITS:
        if rewrite_path(dataset, split, method, condition).exists():
            for item in load_rewrite(dataset, split, method, condition):
                rewrites[item["claim_id"]] = (item.get("claim") or "").strip()

    out: list[tuple[str, list[str]]] = []
    skipped = {"missing": 0, "identical": 0, "flagged": 0}
    for pair_key, members in pairs.items():
        keep: list[str] = []
        for claim_id in members:
            rewrite = rewrites.get(claim_id, "")
            original = base[claim_id]["claim"]
            if not rewrite:
                skipped["missing"] += 1
                continue
            if rewrite == original:
                skipped["identical"] += 1
                continue
            if unflagged_only and flag_reasons(original, rewrite):
                skipped["flagged"] += 1
                continue
            keep.append(claim_id)
        if keep:
            out.append((pair_key, keep))
    return out, skipped


def sample_items(
    dataset: str, n_per_cond: int, rng: random.Random
) -> tuple[list[dict], list[dict]]:
    """Draw n_per_cond pair-disjoint claims per condition, scarcest pool first."""
    base = load_base(dataset)
    pairs = build_pairs(base)

    pools = {}
    for condition, method, type_label, unflagged_only in CONDITIONS:
        avail, skipped = eligible_pairs(
            dataset, condition, method, unflagged_only, base, pairs
        )
        rewrites = {}
        for split in SPLITS:
            if rewrite_path(dataset, split, method, condition).exists():
                for item in load_rewrite(dataset, split, method, condition):
                    rewrites[item["claim_id"]] = (item.get("claim") or "").strip()
        pools[type_label] = {
            "condition": condition,
            "method": method,
            "unflagged_only": unflagged_only,
            "avail": avail,
            "skipped": skipped,
            "rewrites": rewrites,
        }

    items: list[dict] = []
    report: list[dict] = []
    used: set[str] = set()
    # Scarcest pool first so the filtered conditions are not starved by the
    # unconstrained ones competing for the same pairs.
    for type_label in sorted(pools, key=lambda t: len(pools[t]["avail"])):
        pool = pools[type_label]
        candidates = [(p, m) for p, m in pool["avail"] if p not in used]
        rng.shuffle(candidates)
        taken = candidates[:n_per_cond]
        for pair_key, members in taken:
            used.add(pair_key)
            claim_id = rng.choice(members)
            info = base[claim_id]
            items.append({
                "type": type_label,
                "condition": pool["condition"],
                "method": pool["method"],
                "unflagged_only": pool["unflagged_only"],
                "claim_id": claim_id,
                "claim_id_pair": info["claim_id_pair"],
                "split": info["split"],
                "label": info["label"],
                "evi_path": info["evi_path"],
                "original_claim": info["claim"],
                "rewritten_claim": pool["rewrites"][claim_id],
            })
        report.append({
            "type": type_label,
            "condition": pool["condition"],
            "unflagged_only": pool["unflagged_only"],
            "pool": len(pool["avail"]),
            "pool_disjoint": len(candidates),
            "taken": len(taken),
            "skipped": pool["skipped"],
        })

    if len(used) != len(items):
        raise RuntimeError("sampled claims are not pair-disjoint")
    return items, report


def assign_annotator_pairs(
    items: list[dict],
) -> tuple[dict[str, tuple[str, str]], dict[str, int], dict[tuple[str, str], int]]:
    """
    Give every item two annotators.

    Greedy on (uses within this condition, uses overall, current annotator load).
    Leading with per-condition usage spreads all 15 annotator pairs evenly, so
    every pairwise agreement cell is populated to comparable precision; the load
    term then balances per-annotator workload to within one item.
    """
    pair_list = list(combinations(ANNOTATORS, 2))
    loads = {a: 0 for a in ANNOTATORS}
    total_use = {p: 0 for p in pair_list}
    assigned: dict[str, tuple[str, str]] = {}

    for type_label in dict.fromkeys(item["type"] for item in items):
        cond_use = {p: 0 for p in pair_list}
        for item in [i for i in items if i["type"] == type_label]:
            chosen = min(
                pair_list,
                key=lambda p: (cond_use[p], total_use[p], loads[p[0]] + loads[p[1]]),
            )
            assigned[item["item_id"]] = chosen
            loads[chosen[0]] += 1
            loads[chosen[1]] += 1
            cond_use[chosen] += 1
            total_use[chosen] += 1

    spread = max(loads.values()) - min(loads.values())
    if spread > 1:
        raise RuntimeError(f"annotator loads differ by {spread}")
    return assigned, loads, total_use


def write_workbook(
    items: list[dict],
    assigned: dict[str, tuple[str, str]],
    out_path: Path,
    rng: random.Random,
    mark_diff: bool = True,
) -> None:
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    plain = InlineFont()
    highlight = InlineFont(b=True, color="FFC00000")

    def rich(runs: list[tuple[str, bool]]):
        """Rich-text cell value, or a plain string when nothing is marked."""
        if not any(marked for _, marked in runs):
            return "".join(segment for segment, _ in runs)
        return CellRichText([
            TextBlock(highlight if marked else plain, segment)
            for segment, marked in runs
        ])

    by_annotator: dict[str, list[dict]] = {a: [] for a in ANNOTATORS}
    for item in items:
        for annotator in assigned[item["item_id"]]:
            by_annotator[annotator].append(item)

    wb = Workbook()
    wb.remove(wb.active)
    for annotator in ANNOTATORS:
        rows = by_annotator[annotator][:]
        rng.shuffle(rows)
        ws = wb.create_sheet(annotator)
        ws.append(SHEET_COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        if mark_diff:
            header = ws.cell(1, SHEET_COLUMNS.index("rewritten_claim") + 1)
            header.comment = Comment(DIFF_LEGEND, "build_annotation_sheets")

        for row in rows:
            ws.append([row["item_id"], row["original_claim"], "", "",
                       row["evi_path"], "", ""])
            cell = ws.cell(ws.max_row, SHEET_COLUMNS.index("rewritten_claim") + 1)
            cell.value = (
                rich(diff_runs(row["original_claim"], row["rewritten_claim"]))
                if mark_diff
                else row["rewritten_claim"]
            )

        for idx, (name, width) in enumerate(zip(SHEET_COLUMNS, COLUMN_WIDTHS), 1):
            letter = get_column_letter(idx)
            ws.column_dimensions[letter].width = width
            wrap = name in WRAPPED
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=wrap, vertical="top")
        ws.freeze_panes = "A2"

        col = get_column_letter(SHEET_COLUMNS.index("same_as_original") + 1)
        dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
        dv.prompt = "TRUE if the rewrite preserves the original meaning, else FALSE"
        dv.promptTitle = "same_as_original"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")

    wb.save(out_path)


def write_master_key(
    items: list[dict], assigned: dict[str, tuple[str, str]], out_path: Path
) -> None:
    with out_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=MASTER_COLUMNS)
        writer.writeheader()
        for item in sorted(items, key=lambda i: i["item_id"]):
            a1, a2 = assigned[item["item_id"]]
            writer.writerow({
                "item_id": item["item_id"],
                "type": item["type"],
                "condition": item["condition"],
                "method": item["method"],
                "split": item["split"],
                "claim_id": item["claim_id"],
                "claim_id_pair": item["claim_id_pair"],
                "label": item["label"],
                "evi_path": item["evi_path"],
                "annotator_1": a1,
                "annotator_2": a2,
                "unflagged_only": item["unflagged_only"],
                "original_claim": item["original_claim"],
                "rewritten_claim": item["rewritten_claim"],
            })


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--dataset", default="sciclaimeval")
    ap.add_argument("--n-per-cond", type=int, default=80)
    ap.add_argument("--seed", type=int, default=20260915)
    ap.add_argument("--out-dir", type=Path, default=OUT_DIR)
    ap.add_argument("--no-diff-marks", action="store_true",
                    help="write claims as plain text instead of marking the word-level diff")
    args = ap.parse_args()

    rng = random.Random(args.seed)
    items, report = sample_items(args.dataset, args.n_per_cond, rng)

    rng.shuffle(items)
    for i, item in enumerate(items, 1):
        item["item_id"] = f"IT-{i:04d}"

    assigned, loads, total_use = assign_annotator_pairs(items)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = args.out_dir / "annotation_assignment.xlsx"
    key_path = args.out_dir / "master_key.csv"
    write_workbook(items, assigned, xlsx_path, random.Random(args.seed + 1),
                   mark_diff=not args.no_diff_marks)
    write_master_key(items, assigned, key_path)

    print(f"dataset={args.dataset} splits={'+'.join(SPLITS)} seed={args.seed}\n")
    header = f"{'type':<5} {'condition':<10} {'sampling':<16} {'pool':>6} {'disjoint':>9} {'taken':>6}"
    print(header)
    print("-" * len(header))
    for row in report:
        mode = "unflagged only" if row["unflagged_only"] else "random"
        print(f"{row['type']:<5} {row['condition']:<10} {mode:<16} "
              f"{row['pool']:>6} {row['pool_disjoint']:>9} {row['taken']:>6}")
        s = row["skipped"]
        print(f"{'':<5} excluded: {s['missing']} missing, {s['identical']} identical, "
              f"{s['flagged']} flagged")

    print(f"\nclaims={len(items)}  judgments={2 * len(items)}")
    print("load: " + "  ".join(f"{a}={loads[a]}" for a in ANNOTATORS))
    used_pairs = sum(1 for v in total_use.values() if v)
    print(f"annotator pairs used: {used_pairs}/{len(total_use)}  "
          f"per-pair items {min(total_use.values())}-{max(total_use.values())}")
    print(f"\nwrote {xlsx_path}")
    print(f"wrote {key_path}")


if __name__ == "__main__":
    main()
