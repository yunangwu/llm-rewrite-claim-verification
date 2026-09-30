#!/usr/bin/env python3
"""
Score the returned human annotations.

Reads ``result/annotation_assignment-<ID>.xlsx`` (one file per annotator, ID =
A1..A6), takes the ``same_as_original`` judgments from that annotator's own
sheet, joins them to ``master_key.csv``, and reports:

  - coverage (claims with two / one / zero judgments)
  - per-condition preservation rate, observed agreement P_o, Fleiss' kappa and
    Gwet's AC1 (agreement measures on doubly judged claims only)
  - per-annotator share of judgments marking a meaning change

Fleiss' kappa is used rather than Cohen's because different annotator pairs
judged different claims. Both kappa and AC1 use the pooled label prevalence of
the doubly judged claims in the row.

Usage:
  python annotation/score_annotations.py
  python annotation/score_annotations.py --latex
"""

from __future__ import annotations

import argparse
import collections
import csv
import re
from pathlib import Path

HERE = Path(__file__).resolve().parent
CONDITIONS = ("CT", "C1", "OA", "D1")


def _to_bool(value) -> bool | None:
    if isinstance(value, bool):
        return value
    text = str(value).strip().upper() if value is not None else ""
    return {"TRUE": True, "FALSE": False}.get(text)


def load_judgments(result_dir: Path, key: dict[str, dict]) -> dict[str, dict[str, bool]]:
    """item_id -> {annotator: judgment}, validated against the master key."""
    from openpyxl import load_workbook

    judgments: dict[str, dict[str, bool]] = collections.defaultdict(dict)
    files = sorted(result_dir.glob("annotation_assignment-A*.xlsx"))
    if not files:
        raise SystemExit(f"no annotation_assignment-A*.xlsx files in {result_dir}")
    for path in files:
        annotator = re.search(r"-(A\d+)$", path.stem).group(1)
        ws = load_workbook(path, data_only=True)[annotator]
        header = [c.value for c in ws[1]]
        col_id = header.index("item_id") + 1
        col_label = header.index("same_as_original") + 1
        for row in range(2, ws.max_row + 1):
            item_id = ws.cell(row, col_id).value
            label = _to_bool(ws.cell(row, col_label).value)
            if item_id is None or label is None:
                continue
            owners = {key[item_id]["annotator_1"], key[item_id]["annotator_2"]}
            if annotator not in owners:
                raise SystemExit(f"{path.name}: {annotator} is not assigned {item_id}")
            judgments[item_id][annotator] = label
    return judgments


def agreement(pairs: list[tuple[bool, bool]]) -> dict[str, float]:
    """P_o, Fleiss' kappa and Gwet's AC1 for claims with exactly two judgments."""
    n = len(pairs)
    p_o = sum(a == b for a, b in pairs) / n
    prev = sum(a + b for a, b in pairs) / (2 * n)
    pe_kappa = prev ** 2 + (1 - prev) ** 2
    pe_ac1 = 2 * prev * (1 - prev)
    kappa = (p_o - pe_kappa) / (1 - pe_kappa) if pe_kappa < 1 else float("nan")
    ac1 = (p_o - pe_ac1) / (1 - pe_ac1) if pe_ac1 < 1 else float("nan")
    return {"n2": n, "p_o": p_o, "prev": prev, "kappa": kappa, "ac1": ac1}


def summarize(key: dict[str, dict], judgments: dict[str, dict[str, bool]], items: list[str]) -> dict:
    labels = [v for i in items for v in judgments.get(i, {}).values()]
    pairs = [tuple(judgments[i].values()) for i in items if len(judgments.get(i, {})) == 2]
    row = {"judg": len(labels), "pres": 100 * sum(labels) / len(labels)}
    row.update(agreement(pairs))
    return row


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--result-dir", type=Path, default=HERE / "result")
    ap.add_argument("--master-key", type=Path, default=HERE / "master_key.csv")
    ap.add_argument("--latex", action="store_true", help="also print LaTeX table rows")
    args = ap.parse_args()

    with args.master_key.open(encoding="utf-8") as f:
        key = {r["item_id"]: r for r in csv.DictReader(f)}
    judgments = load_judgments(args.result_dir, key)

    coverage = collections.Counter(len(judgments.get(i, {})) for i in key)
    print(f"annotators: {sorted({a for d in judgments.values() for a in d})}")
    print(f"claims: {len(key)}  two judgments: {coverage[2]}  one: {coverage[1]}  "
          f"none: {coverage[0]}  judgments: {sum(len(d) for d in judgments.values())}\n")

    rows = {c: summarize(key, judgments, [i for i in key if key[i]["type"] == c])
            for c in CONDITIONS}
    rows["All"] = summarize(key, judgments, list(key))

    print(f"{'Cond.':<5} {'Judg.':>6} {'Pres.%':>7} {'n2':>4} {'P_o':>6} {'kappa':>7} {'AC1':>6}")
    for name, r in rows.items():
        print(f"{name:<5} {r['judg']:>6} {r['pres']:>7.1f} {r['n2']:>4} "
              f"{r['p_o']:>6.3f} {r['kappa']:>7.3f} {r['ac1']:>6.3f}")
    print(f"\nlabel prevalence (preserved) on doubly judged claims: {100 * rows['All']['prev']:.1f}%")

    print("\nper annotator, share of judgments marking a meaning change:")
    for annotator in sorted({a for d in judgments.values() for a in d}):
        own = [d[annotator] for d in judgments.values() if annotator in d]
        print(f"  {annotator}: {len(own):>3} judgments, "
              f"{100 * sum(not v for v in own) / len(own):.1f}% changed")

    if args.latex:
        print("\n% LaTeX rows: Cond. & Judg. & Pres. & n_2 & P_o & kappa & AC1")
        for name, r in rows.items():
            if name == "All":
                print("\\midrule")
            print(f"{name:<3} & {r['judg']:>3} & {r['pres']:.1f} & {r['n2']:>3} & "
                  f"{r['p_o']:.3f} & {r['kappa']:.3f} & {r['ac1']:.3f} \\\\")


if __name__ == "__main__":
    main()
